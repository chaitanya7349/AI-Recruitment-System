from pathlib import Path

from docx import Document
from reportlab.pdfgen import canvas

from app.services.resume_parser import (
    extract_docx,
    extract_pdf,
    extract_text,
)


def test_extract_text_returns_empty_for_unsupported_file():
    result = extract_text("resume.txt")

    assert result == ""


def test_extract_docx():
    file_path = Path("test_resume_parser.docx")

    try:
        document = Document()
        document.add_paragraph("Rahul Kumar")
        document.add_paragraph("Python Developer")
        document.add_paragraph("Python SQL FastAPI")

        document.save(file_path)

        text = extract_docx(str(file_path))

        assert "Rahul Kumar" in text
        assert "Python Developer" in text
        assert "Python SQL FastAPI" in text

    finally:
        file_path.unlink(missing_ok=True)


def test_extract_text_from_docx():
    file_path = Path("test_resume_parser_text.docx")

    try:
        document = Document()
        document.add_paragraph("Candidate Resume")
        document.add_paragraph("Software Engineer")

        document.save(file_path)

        text = extract_text(str(file_path))

        assert "Candidate Resume" in text
        assert "Software Engineer" in text

    finally:
        file_path.unlink(missing_ok=True)


def test_extract_pdf():
    file_path = Path("test_resume_parser.pdf")

    try:
        pdf = canvas.Canvas(str(file_path))
        pdf.drawString(100, 750, "Rahul Kumar")
        pdf.drawString(100, 730, "Python Developer")
        pdf.save()

        text = extract_pdf(str(file_path))

        assert "Rahul Kumar" in text
        assert "Python Developer" in text

    finally:
        file_path.unlink(missing_ok=True)


def test_extract_text_from_pdf():
    file_path = Path("test_resume_parser_text.pdf")

    try:
        pdf = canvas.Canvas(str(file_path))
        pdf.drawString(100, 750, "Candidate Resume")
        pdf.drawString(100, 730, "Backend Engineer")
        pdf.save()

        text = extract_text(str(file_path))

        assert "Candidate Resume" in text
        assert "Backend Engineer" in text

    finally:
        file_path.unlink(missing_ok=True)
